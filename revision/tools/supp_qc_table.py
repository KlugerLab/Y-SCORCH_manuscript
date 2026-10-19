#!/usr/bin/env python3
"""
Supplementary sample-QC table (SCORCH_sample_qc.xlsx) from the CellBender-processed
pipeline, and .xlsx copies of the regenerated DEG tables.

Row definitions follow the published table:
  Cell number                     nuclei called by CellBender
  Cell number after filtering     nuclei after the nFeature, mitochondrial/sex-gene and
                                  doublet filters (the cell-type labelling stage)
  % > 1000 transcripts, % > 500 genes, median transcripts, median genes
                                  computed on those nuclei from the CellBender-corrected counts
  Median mitochondrial %          median percentage of mitochondrial reads in those nuclei,
                                  from the corrected counts
  Nuclei in final analysis        nuclei of the sample in the integrated object analysed in the
                                  paper (after the global filters and the VST sample exclusions)
  HIV detected                    at least one retained nucleus with >= 3 HIV-aligned reads in
                                  snRNA or snATAC
  Sample kept for final analysis  the sample contributes nuclei to the final object

Runs on the host (needs openpyxl): python3 revision/tools/supp_qc_table.py
Inputs: PIPELINE_DIR, FIXED_META_DIR, PUB_SUPP_DIR (layout of the published sheet),
OUT_ROOT/01-Regional_data_processing/combined/data/combined_v3_cell_metadata.csv,
OUT_ROOT/07-Viral_Detection/viral_nuclei_cellbender_{RNA,ATAC}.csv.
Outputs: OUT_ROOT/tables/supplementary_tables/.
"""
import csv, glob, os, statistics as st, sys
from collections import Counter, defaultdict
import openpyxl
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "config"))
from local_paths import P
from openpyxl.styles import Font

PIPE = P["PIPELINE_DIR"]
FIXED = P["FIXED_META_DIR"]
OUT_ROOT = P["OUT_ROOT"]
REPAIRED = f"{OUT_ROOT}/meta/fixed_meta_repaired"
PUB = os.path.join(P["PUB_SUPP_DIR"], "SCORCH_sample_qc.xlsx")
OUT = f"{OUT_ROOT}/tables/supplementary_tables"
os.makedirs(OUT, exist_ok=True)
REGIONS = ["INS", "PFC", "VST"]
csv.field_size_limit(sys.maxsize)

def fixed_meta(region, s):
    rel = f"{region.lower()}/{s}/label_transfer/labeled_metadata_with_pt_mt_{s}.csv"
    rp = f"{REPAIRED}/{rel}"
    return rp if os.path.exists(rp) else f"{FIXED}/{rel}"

# ---- final object and viral detection ----------------------------------------------------
final = Counter()
with open(f"{OUT_ROOT}/01-Regional_data_processing/combined/data/combined_v3_cell_metadata.csv") as fh:
    for r in csv.DictReader(fh):
        final[(r["region"], r["sample_ids"].upper())] += 1
hiv = set()
for a in ("RNA", "ATAC"):
    with open(f"{OUT_ROOT}/07-Viral_Detection/viral_nuclei_cellbender_{a}.csv") as fh:
        for r in csv.DictReader(fh):
            if r["highest_threshold_met"] in (">=3", ">=5"):
                hiv.add((r["region"], r["donor"].upper()))

# ---- per-sample metrics --------------------------------------------------------------------
rows = []
for R in REGIONS:
    for d in sorted(glob.glob(f"{PIPE}/{R.lower()}/*_{R}")):
        s = os.path.basename(d); donor = s.split("_")[0]; cond = "PWoH" if "_CTR_" in s else "PWH"
        rec = dict(region=R, donor=donor, condition=cond, sample=s)
        with open(f"{d}/cellbender_qc/cellbender_qc_summary_{s}.csv") as fh:
            rec["cells_called"] = int(next(csv.DictReader(fh))["n_cells"])
        with open(f"{d}/label_transfer/labeled_metadata_{s}.csv") as fh:
            lab = list(csv.DictReader(fh))
        nc = [float(x["nCount_RNA"]) for x in lab]; nf = [float(x["nFeature_RNA"]) for x in lab]
        rec.update(cells_after_filtering=len(lab),
                   pct_gt1000_transcripts=round(100 * sum(v > 1000 for v in nc) / len(nc), 2),
                   pct_gt500_genes=round(100 * sum(v > 500 for v in nf) / len(nf), 2),
                   median_transcripts=st.median(nc), median_genes=st.median(nf))
        fm = fixed_meta(R, s)
        with open(fm) as fh:
            mt = [float(x["percent.mt"]) for x in csv.DictReader(fh)]
        rec["median_pct_mito"] = round(st.median(mt), 4)
        rec["nuclei_final_analysis"] = final.get((R, donor.upper()), 0)
        rec["HIV_detected"] = "Yes" if (R, donor.upper()) in hiv else "No"
        rec["kept_final_analysis"] = "Yes" if rec["nuclei_final_analysis"] > 0 else "No"
        rows.append(rec)
        print(f"[qc] {s}: called {rec['cells_called']}, filtered {rec['cells_after_filtering']}, "
              f"final {rec['nuclei_final_analysis']}, HIV {rec['HIV_detected']}")

COLS = ["region", "donor", "condition", "sample", "cells_called", "cells_after_filtering",
        "pct_gt1000_transcripts", "pct_gt500_genes", "median_transcripts", "median_genes",
        "median_pct_mito", "nuclei_final_analysis", "HIV_detected", "kept_final_analysis"]
with open(f"{OUT}/SCORCH_sample_qc_long.csv", "w", newline="") as fh:
    w = csv.DictWriter(fh, fieldnames=COLS); w.writeheader(); w.writerows(rows)

# ---- published donor order ---------------------------------------------------------------
pub = list(openpyxl.load_workbook(PUB, read_only=True).worksheets[0].iter_rows(values_only=True))
order = {}
for r in pub:
    if r[0] in REGIONS and r[1] is not None:
        order.setdefault(r[0], []).extend(str(x) for x in r[1:] if x is not None)

LABELS = [("Cell number (CellBender-called)", "cells_called"),
          ("Cell number after filtering", "cells_after_filtering"),
          ("% Cells with > 1000 Transcripts", "pct_gt1000_transcripts"),
          ("% Cells with > 500 genes", "pct_gt500_genes"),
          ("Median # Transcripts", "median_transcripts"),
          ("Median # Genes", "median_genes"),
          ("Median mitochondrial %", "median_pct_mito"),
          ("Nuclei in final analysis", "nuclei_final_analysis"),
          ("HIV detected", "HIV_detected"),
          ("Sample kept for final analysis", "kept_final_analysis")]
by = {(r["region"], r["donor"].upper()): r for r in rows}
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "sample_qc"
bold = Font(bold=True)
for cond in ("PWH", "PWoH"):
    for R in REGIONS:
        have = [r["donor"] for r in rows if r["region"] == R and r["condition"] == cond]
        # a donor's condition from the pipeline; published donors with no library in this
        # region (10280 INS, 1063 PFC) are PWH - all PWoH ids start with "HC"
        cond_of = lambda d: by[(R, d.upper())]["condition"] if (R, d.upper()) in by else \
            ("PWoH" if d.upper().startswith("HC") else "PWH")
        pub_cond = [d for d in order.get(R, []) if cond_of(d) == cond]
        seq = pub_cond + [h for h in have if h.upper() not in {p.upper() for p in pub_cond}]
        ws.append([cond] + list(range(1, len(seq) + 1))); ws.cell(ws.max_row, 1).font = bold
        ws.append([R] + seq); ws.cell(ws.max_row, 1).font = bold
        for lab, key in LABELS:
            ws.append([lab] + [by[(R, d.upper())][key] if (R, d.upper()) in by else
                               ("No" if key == "kept_final_analysis" else "NA") for d in seq])
        ws.append([])
ws.column_dimensions["A"].width = 34
notes = wb.create_sheet("definitions")
for line in __doc__.strip().splitlines():
    notes.append([line])
wb.save(f"{OUT}/SCORCH_sample_qc.xlsx")
print(f"[qc] wrote {OUT}/SCORCH_sample_qc.xlsx and SCORCH_sample_qc_long.csv ({len(rows)} samples)")

# ---- xlsx copies of the DEG tables ----------------------------------------------------------
for f in sorted(glob.glob(f"{OUT}/DEG_*.csv")):
    if f.endswith("summary.csv"):
        continue
    wb = openpyxl.Workbook(write_only=True); ws = wb.create_sheet(os.path.basename(f)[:-4][:31])
    with open(f) as fh:
        for i, r in enumerate(csv.reader(fh)):
            if i:
                r = [(float(v) if v not in ("", "NA") and v.replace(".", "", 1).replace("-", "", 1).replace("e", "", 1).replace("E", "", 1).isdigit() else v) for v in r]
            ws.append(r)
    wb.save(f[:-4] + ".xlsx")
    print(f"[qc] xlsx: {os.path.basename(f)[:-4]}.xlsx")
