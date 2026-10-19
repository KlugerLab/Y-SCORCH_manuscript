"""Write the summarised cohort table (revision/tools/table1_cohort.R) as a formatted .xlsx.

The R container has no Excel writer, so this runs on the host after the R tool:
    python3 revision/tools/table1_to_xlsx.py
Reads  OUT_ROOT/tables/table1_cohort/table1_cohort.csv and table1_cohort.md (footnote)
Writes OUT_ROOT/tables/table1_cohort/table1_cohort.xlsx
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

import os, sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "config"))
from local_paths import P
D = Path(P["OUT_ROOT"]) / "tables" / "table1_cohort"
rows = list(csv.DictReader(open(D / "table1_cohort.csv", encoding="utf-8")))
cols = list(rows[0].keys())            # section, variable, PWoH (n = ..), PWH (n = ..), p, note
footnote = (D / "table1_cohort.md").read_text(encoding="utf-8").strip().splitlines()[-1]

wb = Workbook()
ws = wb.active
ws.title = "Table 1"
thin = Side(style="thin")
ws.append(["Table 1. Characteristics of the donors"])
ws["A1"].font = Font(bold=True, size=12)
ws.append(["Characteristic", cols[2], cols[3], "P value", "Note"])
for c in ws[2]:
    c.font = Font(bold=True)
    c.border = Border(top=thin, bottom=thin)
section = None
for r in rows:
    if r["section"] != section:
        section = r["section"]
        ws.append([section])
        ws.cell(ws.max_row, 1).font = Font(bold=True, italic=True)
    sub = r["variable"].startswith("   ")
    ws.append([r["variable"].strip(), r[cols[2]], r[cols[3]], r["p"], r["note"]])
    ws.cell(ws.max_row, 1).alignment = Alignment(indent=2 if sub else 1)
    for j in (2, 3, 4):
        ws.cell(ws.max_row, j).alignment = Alignment(horizontal="center")
for c in ws[ws.max_row]:
    c.border = Border(bottom=thin)
ws.append([])
ws.append([footnote])
fr = ws.max_row
ws.merge_cells(start_row=fr, start_column=1, end_row=fr, end_column=5)
ws.cell(fr, 1).alignment = Alignment(wrap_text=True, vertical="top")
ws.cell(fr, 1).font = Font(size=9)
ws.row_dimensions[fr].height = 75
for col, w in zip("ABCDE", (58, 30, 30, 9, 50)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A3"
out = D / "table1_cohort.xlsx"
wb.save(out)
print(out)
